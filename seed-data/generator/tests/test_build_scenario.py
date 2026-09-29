import importlib.util
import json
import time
from pathlib import Path

import pytest
import yaml
from jsonschema import ValidationError
from openpyxl import load_workbook


SCRIPT = Path(__file__).parents[1] / "build_scenario.py"


def load_builder():
    spec = importlib.util.spec_from_file_location("build_scenario", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def ford_pack():
    return {
        "slug": "ford-auto",
        "customer": {"name": "福特中国", "industry": "汽车", "locale": "zh-CN"},
        "course": {"code": "MS-4018", "purpose": "Demo"},
        "tenantDomain": "moneyyu.com",
        "roles": {
            "Admin": {"upn": "admin@moneyyu.com", "displayName": "Admin"},
            "Planner": {"upn": "AdeleV@moneyyu.com", "displayName": "Adele Vance"},
            "Sales": {"upn": "NestorW@moneyyu.com", "displayName": "Nestor Wilke"},
        },
        "teams": {
            "channels": [
                {
                    "channelName": "新车上市",
                    "channelDescription": "工作动态",
                    "messages": [
                        {
                            "fromRole": "Planner",
                            "dayOffset": -2,
                            "time": "09:00",
                            "bodyHtml": "<p>经销商库存复盘</p>",
                            "mentions": [],
                            "replies": [],
                        }
                    ],
                }
            ]
        },
        "emails": {
            "emailThreads": [
                {
                    "threadName": "复盘",
                    "emails": [
                        {
                            "action": "send",
                            "fromRole": "Planner",
                            "toRoles": ["Sales"],
                            "ccRoles": ["Admin"],
                            "subject": "20260929 经销商库存复盘",
                            "bodyHtml": "<p>请查看库存</p>",
                        },
                        {
                            "action": "reply",
                            "fromRole": "Sales",
                            "toRoles": ["Planner"],
                            "ccRoles": ["Admin"],
                            "subject": "Re: 20260929 经销商库存复盘",
                            "bodyHtml": "<p>收到</p>",
                        },
                    ],
                }
            ]
        },
        "workbooks": [
            {
                "filename": "经销商库存.xlsx",
                "sheets": [
                    {
                        "name": "库存",
                        "table": "DealerStock",
                        "headers": ["经销商", "现车", "月销", "库存天数"],
                        "rows": [["北京一店", 120, 30, None]],
                        "formulas": {"库存天数": '=IF(C{row}=0,0,B{row}/C{row}*30)'},
                        "chart": {"category": "经销商", "value": "现车"},
                    }
                ],
            }
        ],
    }


def write_pack(tmp_path, pack):
    path = tmp_path / "pack.yaml"
    path.write_text(yaml.safe_dump(pack, allow_unicode=True), encoding="utf-8")
    return path


def test_ford_outputs_only_declared_surfaces(tmp_path):
    builder = load_builder()
    output = tmp_path / "generated"
    builder.build_scenario(write_pack(tmp_path, ford_pack()), "20260929", output)
    assert {
        "config.json.example",
        "emails.json",
        "teams-messages.json",
        "files-manifest.json",
        "run.ps1",
        "README.md",
    } <= {p.name for p in output.iterdir()}
    assert not (output / "sharepoint-sites.json").exists()
    assert "Invoke-SeedUserProfiles" not in (output / "run.ps1").read_text(encoding="utf-8")
    assert "Invoke-SeedSharePoint" not in (output / "run.ps1").read_text(encoding="utf-8")
    runner = (output / "run.ps1").read_text(encoding="utf-8")
    assert "PreflightOnly" in runner
    assert '$preflight.FilesManifestPath = Join-Path $root "files-manifest.json"' in runner
    assert 'Invoke-UploadFiles.ps1") -ConfigPath $configPath -FilesManifestPath (Join-Path $root "files-manifest.json") -StrictScenario' in runner
    teams = json.loads((output / "teams-messages.json").read_text(encoding="utf-8"))
    assert teams["channels"][0]["messages"][0]["createdDateTime"] == "2026-09-27T09:00:00+08:00"
    book = load_workbook(output / "DEMO-FILE" / "经销商库存.xlsx")
    assert book["库存"]["D2"].value == "=IF(C2=0,0,B2/C2*30)"
    assert "DealerStock" in book["库存"].tables
    assert len(book["库存"]._charts) == 1
    assert "模拟" in book["演示说明"]["A1"].value


def test_deterministic_workbook_bytes(tmp_path):
    builder = load_builder()
    pack = write_pack(tmp_path, ford_pack())
    builder.build_scenario(pack, "20260929", tmp_path / "a")
    time.sleep(1.1)
    builder.build_scenario(pack, "20260929", tmp_path / "b")
    assert (tmp_path / "a" / "DEMO-FILE" / "经销商库存.xlsx").read_bytes() == (
        tmp_path / "b" / "DEMO-FILE" / "经销商库存.xlsx"
    ).read_bytes()


@pytest.mark.parametrize(
    "change",
    [
        lambda p: p["roles"]["Sales"].update(upn="another@example.com"),
        lambda p: p["emails"]["emailThreads"][0]["emails"][1].update(fromRole="Planner"),
        lambda p: p["emails"]["emailThreads"][0]["emails"][0].update(ccRoles=[]),
        lambda p: p["teams"]["channels"][0].update(channelName="新车/上市"),
        lambda p: p["course"].update(purpose="福特 Demo"),
        lambda p: p["emails"]["emailThreads"][0]["emails"][0].update(subject="库存复盘"),
        lambda p: p["roles"].update(New={"createUser": True}),
    ],
)
def test_invalid_pack_fails_without_output(tmp_path, change):
    builder = load_builder()
    pack = ford_pack()
    change(pack)
    output = tmp_path / "generated"
    with pytest.raises((ValueError, ValidationError)):
        builder.build_scenario(write_pack(tmp_path, pack), "20260929", output)
    assert not output.exists()


def test_sharepoint_only_outputs_and_runner(tmp_path):
    builder = load_builder()
    pack = ford_pack()
    del pack["teams"]
    del pack["emails"]
    del pack["workbooks"]
    pack["sharepoint"] = {
        "sites": [
            {
                "alias": "ms4022-productsupport-{{DATE}}",
                "displayName": "MS-4022 - Product support - {{DATE}}",
                "owners": ["Admin"],
                "members": ["Sales"],
                "documents": [],
                "lists": [],
                "documentLibrary": "Products",
            }
        ]
    }
    output = tmp_path / "generated"
    builder.build_scenario(write_pack(tmp_path, pack), "20260929", output)
    assert (output / "sharepoint-sites.json").exists()
    assert not (output / "teams-messages.json").exists()
    assert not (output / "emails.json").exists()
    assert not (output / "files-manifest.json").exists()
    runner = (output / "run.ps1").read_text(encoding="utf-8")
    assert "Invoke-SeedSharePoint" in runner
    assert "Invoke-SeedEmails" not in runner
    assert "Invoke-UploadFiles" not in runner


def test_bad_generation_rule_cannot_leave_partial_scenario(tmp_path):
    builder = load_builder()
    pack = ford_pack()
    sheet = pack["workbooks"][0]["sheets"][0]
    sheet["headers"] = ["经销商", "月份", "销量"]
    sheet.pop("rows")
    sheet.pop("formulas")
    sheet["generate"] = {
        "seed": 42,
        "dimensions": ["北京"],
        "periods": ["9月"],
        "min": 10,
        "max": 50,
        "overrides": [{"dimension": "上海", "period": "9月", "value": 100}],
    }
    output = tmp_path / "invalid"
    with pytest.raises(ValueError, match="Override"):
        builder.build_scenario(write_pack(tmp_path, pack), "20260929", output)
    assert not output.exists()


def test_ford_pack_has_consistent_simulated_facts():
    builder = load_builder()
    pack = yaml.safe_load(
        (SCRIPT.parents[1] / "packs" / "ms4018-ford-auto" / "pack.yaml").read_text(
            encoding="utf-8"
        )
    )
    builder.validate_pack(pack, "20260929")
    assert len(pack["teams"]["channels"]) == 5
    assert len(pack["emails"]["emailThreads"]) == 5
    assert len(pack["workbooks"]) == 5
    sheets = {
        book["filename"]: book["sheets"] for book in pack["workbooks"]
    }
    assert sum(r[2] for r in sheets["新车销售追踪.xlsx"][0]["rows"]) == 860
    assert sum(r[3] for r in sheets["新车销售追踪.xlsx"][0]["rows"]) == 960
    monthly = sheets["新车销售追踪.xlsx"][0]["rows"]
    regional = sheets["新车销售追踪.xlsx"][1]["rows"]
    for month, _, actual, target, _ in monthly:
        assert sum(r[2] for r in regional if r[0] == month) == actual
        assert sum(r[3] for r in regional if r[0] == month) == target
    assert sheets["经销商库存.xlsx"][0]["rows"][0][:3] == ["上海一店", 240, 30]
    assert sheets["零部件品质追踪.xlsx"][0]["rows"][0][2:4] == [1800, 1470]
    assert sheets["零部件品质追踪.xlsx"][1]["rows"][0][2:4] == [600, 36]
    assert sheets["售后NPS.xlsx"][0]["rows"][0][1:4] == [62, 31, 100]
    assert sheets["工厂KPI.xlsx"][0]["rows"][0][1:3] == [1000, 924]


def test_ms4022_archive_requires_all_documents_before_graph_write(tmp_path):
    builder = load_builder()
    pack = yaml.safe_load(
        (SCRIPT.parents[1] / "packs" / "ms4022-productsupport" / "pack.yaml").read_text(
            encoding="utf-8"
        )
    )
    output = tmp_path / "ms4022"
    builder.build_scenario(write_pack(tmp_path, pack), "20260929", output)
    runner = (output / "run.ps1").read_text(encoding="utf-8")
    assert "-ne 9" in runner
    assert runner.index("-ne 9") < runner.index("Invoke-SeedSharePoint")
    assert "GetRelativePath($source, $_.FullName)" in runner
    assert "$config.filesSourceDir" in runner
    assert "[System.IO.Path]::IsPathRooted($config.filesSourceDir)" in runner
    assert "StartsWith($rootFull" in runner
    assert "Join-Path (Split-Path $source -Parent) \"Products.zip\"" in runner


def test_bad_table_name_is_rejected_before_scenario_creation(tmp_path):
    builder = load_builder()
    pack = ford_pack()
    pack["workbooks"][0]["sheets"][0]["table"] = "Invalid Table"
    output = tmp_path / "invalid-table"
    with pytest.raises(ValueError, match="table"):
        builder.build_scenario(write_pack(tmp_path, pack), "20260929", output)
    assert not output.exists()
