# /// script
# requires-python = ">=3.10"
# dependencies = ["pyyaml", "jsonschema", "openpyxl"]
# ///
"""Generate a dated, customer-neutral Microsoft 365 demo scenario from a YAML pack."""

import argparse
import io
import json
import random
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import yaml
from jsonschema import validate
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HERE = Path(__file__).resolve().parent
LOCAL_ZONE = timezone(timedelta(hours=8))


def check_roles(pack, roles):
    for role in roles:
        if role not in pack["roles"]:
            raise ValueError(f"Undeclared role: {role}")


def validate_pack(pack, date):
    with (HERE / "pack.schema.json").open(encoding="utf-8") as source:
        validate(pack, json.load(source))
    datetime.strptime(date, "%Y%m%d")
    if pack["customer"]["name"].casefold() in pack["course"]["purpose"].casefold():
        raise ValueError("Customer brand must not appear in persistent resource names")
    upns = []
    for role in pack["roles"].values():
        upn = role["upn"]
        if not upn.lower().endswith("@" + pack["tenantDomain"].lower()):
            raise ValueError("All roles must use the declared tenant domain")
        upns.append(upn.casefold())
    if len(upns) != len(set(upns)):
        raise ValueError("Each role must use a distinct existing account")
    if "teams" in pack:
        names = set()
        for channel in pack["teams"]["channels"]:
            name = channel["channelName"]
            if re.search(r'[+#%&\\/:<>?|\"]', name) or name in names:
                raise ValueError(f"Invalid or repeated channel name: {name}")
            names.add(name)
            for message in channel["messages"]:
                for item in (message, *message.get("replies", [])):
                    check_roles(pack, [item["fromRole"], *item.get("mentions", [])])
                    if int(item["dayOffset"]) > 0:
                        raise ValueError("Future-dated migration messages are not supported")
    if "emails" in pack:
        subjects = set()
        for thread in pack["emails"]["emailThreads"]:
            emails = thread["emails"]
            if not emails or emails[0]["action"] != "send":
                raise ValueError("A thread must start with send")
            first = emails[0]["subject"]
            if date not in first or first in subjects:
                raise ValueError("First email subject must be unique and contain the demo date")
            subjects.add(first)
            for i, email in enumerate(emails):
                check_roles(pack, [email["fromRole"], *email["toRoles"], *email["ccRoles"]])
                if "Admin" not in email["ccRoles"]:
                    raise ValueError("The observer must be CC'd on every email")
                if i:
                    if email["action"] != "reply" or email["fromRole"] not in (
                        emails[i - 1]["toRoles"] + emails[i - 1]["ccRoles"]
                    ):
                        raise ValueError("Reply author must have received the previous email")
    if "sharepoint" in pack:
        for site in pack["sharepoint"]["sites"]:
            if pack["customer"]["name"].casefold() in (
                site["displayName"] + site["alias"]
            ).casefold():
                raise ValueError("Customer brand must not appear in group identity")
            check_roles(pack, site["owners"] + site["members"])
            if "Admin" not in site["owners"]:
                raise ValueError("Admin must own each SharePoint site")
    for book in pack.get("workbooks", []):
        if Path(book["filename"]).name != book["filename"] or not book["filename"].endswith(".xlsx"):
            raise ValueError("Workbook filename must be a basename ending in .xlsx")
        table_names = set()
        for sheet in book["sheets"]:
            table_name = sheet["table"]
            if not re.fullmatch(r"[A-Za-z_][A-Za-z_0-9]*", table_name) or table_name in table_names:
                raise ValueError(f"Invalid or duplicate Excel table name: {table_name}")
            table_names.add(table_name)
            if len(sheet["name"]) > 31 or re.search(r"[\[\]\\/*?:]", sheet["name"]):
                raise ValueError(f"Invalid Excel sheet name: {sheet['name']}")
            if not sheet.get("rows") and not sheet.get("generate"):
                raise ValueError("Each sheet needs rows or generate rules")
            if len(sheet["headers"]) != len(set(sheet["headers"])):
                raise ValueError("Column headers must be unique")
            if sheet.get("generate"):
                if len(sheet["headers"]) != 3:
                    raise ValueError("Generated sheets need exactly three headers")
                rows_for_sheet(sheet)
            for row in sheet.get("rows", []):
                if len(row) != len(sheet["headers"]):
                    raise ValueError("Row width differs from headers")
            for formula_header in sheet.get("formulas", {}):
                if formula_header not in sheet["headers"]:
                    raise ValueError(f"Formula column does not exist: {formula_header}")
            if "chart" in sheet and any(
                sheet["chart"][key] not in sheet["headers"] for key in ("category", "value")
            ):
                raise ValueError("Chart column does not exist")


def write_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def get_messages(pack, date):
    result = {"channels": []}
    anchor = datetime.strptime(date, "%Y%m%d").replace(tzinfo=LOCAL_ZONE)
    for channel in pack["teams"]["channels"]:
        created = {k: v for k, v in channel.items() if k != "messages"}
        created["messages"] = []
        for index, message in enumerate(channel["messages"], start=1):
            record = dict(message)
            record["order"] = index
            replies = record.pop("replies", [])
            for item in (record, *replies):
                day = item.pop("dayOffset")
                hour, minute = map(int, item.pop("time").split(":"))
                item["createdDateTime"] = (anchor + timedelta(days=day, hours=hour, minutes=minute)).isoformat(
                    timespec="seconds"
                )
            record["replies"] = replies
            created["messages"].append(record)
        result["channels"].append(created)
    return result


def rows_for_sheet(sheet):
    if "generate" not in sheet:
        return sheet["rows"]
    rule = sheet["generate"]
    rng = random.Random(rule["seed"])
    rows = []
    for dimension in rule["dimensions"]:
        for period in rule["periods"]:
            rows.append([dimension, period, rng.randint(rule["min"], rule["max"])])
    for override in rule.get("overrides", []):
        match = [r for r in rows if r[:2] == [override["dimension"], override["period"]]]
        if len(match) != 1:
            raise ValueError("Override must match one generated row")
        match[0][2] = override["value"]
    return rows


def make_workbook(book):
    wb = Workbook()
    wb.properties.creator = "MS-4018 demo"
    wb.properties.created = datetime(2020, 1, 1)
    wb.properties.modified = datetime(2020, 1, 1)
    wb.properties.description = "Fictional training data, not customer-reported results."
    for index, sheet in enumerate(book["sheets"]):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = sheet["name"]
        ws.append(sheet["headers"])
        for row in rows_for_sheet(sheet):
            ws.append(row)
            for header, template in sheet.get("formulas", {}).items():
                column = sheet["headers"].index(header) + 1
                ws.cell(ws.max_row, column, template.replace("{row}", str(ws.max_row)))
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="174A75")
            cell.alignment = Alignment(horizontal="center")
        for cells in ws.iter_rows(min_row=2):
            for cell in cells:
                cell.font = Font(name="Arial", color="000000" if cell.data_type == "f" else "0000FF")
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(
                36, max(13, max(len(str(cell.value or "")) for cell in column) + 4)
            )
        ws.freeze_panes = "A2"
        table = Table(displayName=sheet["table"], ref=f"A1:{ws.cell(ws.max_row, ws.max_column).coordinate}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        if "chart" in sheet:
            category = sheet["headers"].index(sheet["chart"]["category"]) + 1
            value = sheet["headers"].index(sheet["chart"]["value"]) + 1
            chart = BarChart()
            chart.title = sheet["chart"]["value"]
            chart.add_data(Reference(ws, min_col=value, min_row=1, max_row=ws.max_row), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=category, min_row=2, max_row=ws.max_row))
            ws.add_chart(chart, f"A{ws.max_row + 3}")
    notes = wb.create_sheet("演示说明")
    notes["A1"] = "模拟教学数据，非客户官方业绩或已确认的召回信息。"
    notes["A2"] = "来源：本地 Customer Pack；所有结论请根据数据表重新核对。"
    notes.column_dimensions["A"].width = 75
    notes["A1"].font = Font(name="Arial", bold=True, color="9C2F22")
    notes["A2"].font = Font(name="Arial")
    stream = io.BytesIO()
    wb.save(stream)
    output = io.BytesIO()
    with ZipFile(io.BytesIO(stream.getvalue())) as archive, ZipFile(output, "w") as stable:
        for name in archive.namelist():
            info = ZipInfo(name, (2020, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            content = archive.read(name)
            if name == "docProps/core.xml":
                content, count = re.subn(
                    rb"(<dcterms:modified[^>]*>)[^<]*",
                    rb"\g<1>2020-01-01T00:00:00Z",
                    content,
                )
                if count != 1:
                    raise ValueError("Workbook modified timestamp was not found")
            stable.writestr(info, content)
    return output.getvalue()


def runner(pack):
    lines = [
        'param([switch]$PreflightOnly)',
        '$ErrorActionPreference = "Stop"',
        '$root = $PSScriptRoot',
        '$engine = Join-Path $root "..\\..\\engine"',
        '$configPath = Join-Path $root "config.json"',
        'if (-not (Test-Path $configPath)) { throw "Copy config.json.example to config.json first." }',
        '$conn = & (Join-Path $engine "Connect-GraphApp.ps1") -ConfigPath $configPath',
        '$global:AccessToken = $conn.AccessToken',
        'if (-not $global:AccessToken) { throw "No Graph access token returned." }',
        '$preflight = @{ ConfigPath = $configPath }',
    ]
    for key, param, file in (
        ("emails", "EmailsPath", "emails.json"),
        ("teams", "TeamsMessagesPath", "teams-messages.json"),
        ("sharepoint", "SharePointPath", "sharepoint-sites.json"),
        ("workbooks", "FilesManifestPath", "files-manifest.json"),
    ):
        if key in pack:
            lines.append(f'$preflight.{param} = Join-Path $root "{file}"')
    lines += [
        '& (Join-Path $engine "Invoke-SeedPreflight.ps1") @preflight',
        'if ($PreflightOnly) { Write-Host "Read-only preflight complete."; return }',
    ]
    if "workbooks" in pack:
        lines.append('& (Join-Path $engine "Invoke-UploadFiles.ps1") -ConfigPath $configPath -FilesManifestPath (Join-Path $root "files-manifest.json") -StrictScenario')
    if "emails" in pack:
        lines.append('& (Join-Path $engine "Invoke-SeedEmails.ps1") -ConfigPath $configPath -EmailsPath (Join-Path $root "emails.json")')
    if "teams" in pack:
        lines.append('& (Join-Path $engine "Invoke-SeedTeamsChannel.ps1") -ConfigPath $configPath -TeamsMessagesPath (Join-Path $root "teams-messages.json")')
    if "sharepoint" in pack:
        if "sourceArchive" in pack:
            source = pack["sourceArchive"]
            lines += [
                '$config = Get-Content $configPath -Raw | ConvertFrom-Json',
                'if (-not $config.filesSourceDir -or [System.IO.Path]::IsPathRooted($config.filesSourceDir)) { throw "filesSourceDir must be relative to the scenario root." }',
                '$rootFull = [System.IO.Path]::GetFullPath($root)',
                '$source = [System.IO.Path]::GetFullPath((Join-Path $root $config.filesSourceDir))',
                'if (-not $source.StartsWith($rootFull + [System.IO.Path]::DirectorySeparatorChar, [System.StringComparison]::OrdinalIgnoreCase)) { throw "filesSourceDir must stay inside the scenario root." }',
                '$zip = Join-Path (Split-Path $source -Parent) "Products.zip"',
                'if (-not (Test-Path $source)) {',
                '    New-Item -ItemType Directory -Path $source -Force | Out-Null',
                f'    Invoke-WebRequest -Uri "{source["url"]}" -OutFile $zip',
                '    Expand-Archive -Path $zip -DestinationPath $source',
                '}',
            ]
            if source.get("extraFileUrl"):
                lines += [
                    f'$extra = Join-Path $source "{source["extraFileName"]}"',
                    f'if (-not (Test-Path $extra)) {{ Invoke-WebRequest -Uri "{source["extraFileUrl"]}" -OutFile $extra }}',
                ]
            lines += [
                '$sourceFiles = @(Get-ChildItem -Path $source -File -Recurse)',
                f'if ($sourceFiles.Count -ne {source["expectedFileCount"]}) {{ throw "Incomplete lab source: expected {source["expectedFileCount"]} documents; check source directory without deleting files." }}',
                '$siteConfig = Get-Content (Join-Path $root "sharepoint-sites.json") -Raw | ConvertFrom-Json',
                '$siteConfig.sites[0].documents = @($sourceFiles | ForEach-Object {',
                '    @{ sourceFilename = [System.IO.Path]::GetRelativePath($source, $_.FullName); targetFilename = $_.Name }',
                '})',
                'if (-not $siteConfig.sites[0].documents.Count) { throw "No source documents found." }',
                '$siteConfig | ConvertTo-Json -Depth 20 | Set-Content (Join-Path $root "sharepoint-runtime.json") -Encoding utf8',
                '$sharepointPath = Join-Path $root "sharepoint-runtime.json"',
            ]
        else:
            lines.append('$sharepointPath = Join-Path $root "sharepoint-sites.json"')
        lines.append('& (Join-Path $engine "Invoke-SeedSharePoint.ps1") -ConfigPath $configPath -SharePointPath $sharepointPath')
    lines.append('Write-Host "Seeding phases completed. Verify tenant state read-only."')
    return "\n".join(lines) + "\n"


def build_scenario(pack_path, date, output):
    pack_path, output = Path(pack_path), Path(output)
    pack = yaml.safe_load(pack_path.read_text(encoding="utf-8"))
    validate_pack(pack, date)
    if output.exists() and any(output.iterdir()):
        raise FileExistsError(f"Scenario already exists; never overwrite seeded assets: {output}")
    output.mkdir(parents=True, exist_ok=True)
    config = {
        "tenantId": "<tenant-guid>",
        "clientId": "<app-client-id>",
        "clientSecret": "<local-secret>",
        "adminUpn": pack["roles"]["Admin"]["upn"],
        "demoUserUpn": pack["roles"]["Admin"]["upn"],
        "teamDisplayName": f'{pack["course"]["code"]} {pack["course"]["purpose"]} — {date[:4]}-{date[4:6]}-{date[6:]}',
        "teamDescription": f'{pack["course"]["code"]} dated training demo',
        "filesSourceDir": pack.get("sourceArchive", {}).get("dir", "DEMO-FILE"),
        "roles": pack["roles"],
    }
    write_json(output / "config.json.example", config)
    if "teams" in pack:
        write_json(output / "teams-messages.json", get_messages(pack, date))
    if "emails" in pack:
        emails = {"emailThreads": pack["emails"]["emailThreads"]}
        for thread in emails["emailThreads"]:
            for order, email in enumerate(thread["emails"], 1):
                email["order"] = order
        write_json(output / "emails.json", emails)
    if "workbooks" in pack:
        folder = output / "DEMO-FILE"
        folder.mkdir(exist_ok=True)
        for book in pack["workbooks"]:
            (folder / book["filename"]).write_bytes(make_workbook(book))
        write_json(
            output / "files-manifest.json",
            {
                "targetFolder": f'{pack["course"]["code"].replace("-", "")}-{pack["course"]["purpose"]}-{date}',
                "uploadToRole": "Admin",
                "files": [{"localName": book["filename"]} for book in pack["workbooks"]],
            },
        )
    if "sharepoint" in pack:
        sites = json.loads(json.dumps(pack["sharepoint"]).replace("{{DATE}}", date))
        write_json(output / "sharepoint-sites.json", sites)
    (output / "run.ps1").write_text(runner(pack), encoding="utf-8")
    (output / "README.md").write_text(
        f'# {pack["course"]["code"]} {pack["slug"]} ({date})\n\n'
        "Fictional demonstration content; customer names are narrative only. "
        "Only existing, enabled tenant accounts may be used. This seeder never creates users.\n\n"
        "Copy `config.json.example` to locally ignored `config.json`, provide app credentials, "
        "then run `pwsh -File .\\run.ps1 -PreflightOnly` before `pwsh -File .\\run.ps1`. "
        "Stop on any ambiguity; never delete, rename, or repair tenant resources by name. "
        "Outlook messages display execution date (not historical time).\n",
        encoding="utf-8",
    )
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pack", type=Path, required=True)
    parser.add_argument("--date", required=True, help="YYYYMMDD")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(build_scenario(args.pack, args.date, args.output))
